import json
import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

# Helper to load metrics
def load_metrics():
    metrics_path = Path(__file__).parent / "benchmarks" / "retrieval_metrics.json"
    if metrics_path.exists():
        with open(metrics_path, "r") as f:
            return json.load(f)
    return {}

def add_title(doc, text, size=22, bold=True, color=RGBColor(0x1F, 0x49, 0x7D)):
    p = doc.add_heading(level=1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return p

def add_heading(doc, text, level=2, size=16, color=RGBColor(0x2E, 0x5B, 0x8F)):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p

def generate_report():
    metrics = load_metrics()
    doc = Document()
    
    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        
    add_title(doc, "Week 6 Deliverables: Search & Retrieval Optimization")
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.add_run(f"Generated on: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n").italic = True
    p_meta.add_run("Author: AI Engineering Team\nProject: SIA RAG System")
    
    doc.add_paragraph()
    
    add_heading(doc, "1. Executive Summary")
    doc.add_paragraph(
        "This report outlines the evaluation of different search algorithms for the SIA RAG platform. "
        "We benchmarked four distinct retrieval strategies to find the optimal balance between "
        "search quality (Recall & Precision) and latency."
    )
    
    add_heading(doc, "2. Retrieval Strategies Evaluated")
    strategies = [
        ("Sparse Search (BM25)", "Lexical search that matches exact keywords. Fast but misses semantic meaning."),
        ("Dense Search (MiniLM)", "Semantic search using LanceDB and sentence-transformers. Understands intent."),
        ("Hybrid Search (RRF)", "Combines Sparse and Dense search results using Reciprocal Rank Fusion."),
        ("Reranking (Cross-Encoder)", "Passes top Hybrid results through a cross-encoder model to perfectly score query-chunk pairs.")
    ]
    for name, desc in strategies:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(name + ": ").bold = True
        p.add_run(desc)
        
    add_heading(doc, "3. Benchmark Results")
    
    # Add table
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    headers = ["Strategy", "Latency (ms)", "Recall@3", "MRR"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        
    for model_name, data in metrics.items():
        row_cells = table.add_row().cells
        row_cells[0].text = model_name
        row_cells[1].text = f"{data.get('avg_latency_ms', 0):.1f} ms"
        row_cells[2].text = f"{data.get('recall@3', 0):.2f}"
        row_cells[3].text = f"{data.get('mrr', 0):.2f}"
        
    doc.add_paragraph()
    
    add_heading(doc, "4. Performance Charts")
    
    img1 = Path(__file__).parent / "benchmarks" / "search_quality.png"
    if img1.exists():
        doc.add_picture(str(img1), width=Inches(5.5))
        p = doc.add_paragraph("Figure 1: Search Quality Comparison")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    img2 = Path(__file__).parent / "benchmarks" / "search_latency.png"
    if img2.exists():
        doc.add_picture(str(img2), width=Inches(5.5))
        p = doc.add_paragraph("Figure 2: Search Latency Comparison")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    add_heading(doc, "5. Recommendations")
    doc.add_paragraph(
        "1. For maximum speed: Dense Search provides excellent recall with sub-20ms latency.\n"
        "2. For maximum accuracy: The Cross-Encoder Reranker achieves perfect MRR, but incurs a ~400ms latency penalty.\n"
        "3. Final Decision: For production, we recommend continuing with Dense Search (LanceDB) as the default "
        "due to its exceptional speed-to-quality ratio, keeping Reranking as an optional flag for complex analytical queries."
    )
    
    out_path = Path(__file__).parent / "Week6_Search_Optimization_Report.docx"
    doc.save(out_path)
    print(f"Generated {out_path}")

if __name__ == "__main__":
    generate_report()
